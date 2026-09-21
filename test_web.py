from io import BytesIO
from pathlib import Path
import unittest
from openpyxl import Workbook, load_workbook
from app import app
from consolidator import consolidate, ConsolidationError


def workbook(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    data = BytesIO()
    book.save(data)
    data.seek(0)
    return data


class ConsolidatorTests(unittest.TestCase):
    def test_align_and_keep_duplicates(self):
        result = consolidate([
            ('a.xlsx', workbook([['Name', 'Units'], ['Pen', 3], ['Pen', 3]])),
            ('b.xlsx', workbook([['Units', 'Name'], [4, 'Book']]))])
        self.assertEqual(result.rows, 3)
        book = load_workbook(result.content)
        self.assertEqual(list(book.active.values), [('Name', 'Units'), ('Pen', 3), ('Pen', 3), ('Book', 4)])
        book.close()

    def test_bad_inputs(self):
        cases = [
            ('b.xlsx', workbook([['Other'], [1]])),
            ('b.xlsx', BytesIO(b'broken')),
            ('b.csv', BytesIO(b'Name')),
            ('b.xlsx', workbook([['Name'], ['=1+2']])),
            ('b.xlsx', workbook([['Name', 'Name'], [1, 2]])),
            ('b.xlsx', workbook([])),
        ]
        for name, data in cases:
            with self.subTest(name=name, data=data), self.assertRaises(ConsolidationError):
                consolidate([('a.xlsx', workbook([['Name'], ['Pen']])), (name, data)])

    def test_empty_and_blank_rows(self):
        result = consolidate([('a.xlsx', workbook([['Name'], [None], ['Pen']])),
                              ('b.xlsx', workbook([['Name']]))])
        self.assertEqual(result.rows, 1)

    def test_literal_formula_text(self):
        data = workbook([['Name'], ['placeholder']])
        book = load_workbook(data)
        book.active['A2'] = '=literal'
        book.active['A2'].data_type = 's'
        source = BytesIO()
        book.save(source)
        source.seek(0)
        result = consolidate([('a.xlsx', source), ('b.xlsx', workbook([['Name']]))])
        exported = load_workbook(result.content)
        self.assertEqual(exported.active['A2'].data_type, 's')

    def test_web_download_and_errors(self):
        client = app.test_client()
        self.assertEqual(client.get('/').status_code, 200)
        response = client.post('/consolidate', data={'files': [
            (workbook([['Name'], ['Pen']]), 'a.xlsx'),
            (workbook([['Name'], ['Book']]), 'b.xlsx')]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['X-Row-Count'], '2')
        self.assertEqual(response.headers['X-File-Count'], '2')
        self.assertIn('consolidated.xlsx', response.headers['Content-Disposition'])
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.active.max_row, 3)
        self.assertEqual(client.post('/consolidate').status_code, 400)
        old_limit = app.config['MAX_CONTENT_LENGTH']
        try:
            app.config['MAX_CONTENT_LENGTH'] = 10
            self.assertEqual(client.post('/consolidate', data=b'x' * 20, content_type='multipart/form-data').status_code, 413)
        finally:
            app.config['MAX_CONTENT_LENGTH'] = old_limit

    def test_existing_clean_fixture(self):
        folder = Path(__file__).parent / 'tests' / '00_clean' / 'input'
        if not folder.exists():
            self.skipTest('Optional local sample datasets are not present')
        result = consolidate([(path.name, BytesIO(path.read_bytes())) for path in sorted(folder.glob('*.xlsx'))])
        self.assertEqual(result.rows, 600)


if __name__ == '__main__':
    unittest.main()
