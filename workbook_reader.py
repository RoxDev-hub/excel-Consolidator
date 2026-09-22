"""Shared Excel input rules, independent of the web interface."""
from dataclasses import dataclass
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_EXPANDED_BYTES = 100 * 1024 * 1024

class ConsolidationError(ValueError):
    pass

@dataclass
class WorkbookData:
    headers: list
    rows: list
    expanded_bytes: int

def read_workbook(name, stream, budget=MAX_EXPANDED_BYTES):
    if not name.lower().endswith('.xlsx'):
        raise ConsolidationError(f'{name}: use an .xlsx workbook.')
    try:
        with ZipFile(stream) as archive:
            size = sum(item.file_size for item in archive.infolist())
            if size > budget:
                raise ConsolidationError('These workbooks are too large when unpacked. Try smaller files.')
        stream.seek(0)
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ConsolidationError(f'{name}: maximum {MAX_COLUMNS} columns supported.')
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            headers = [cell.value for cell in first]
            if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
                raise ConsolidationError(f'{name}: row 1 must contain a text heading for every column.')
            if len(set(headers)) != len(headers) or any(c.data_type == 'f' for c in first):
                raise ConsolidationError(f'{name}: headings must be unique text, without formulas.')
            rows = []
            for number, row in enumerate(iterator, start=2):
                if number > MAX_ROWS + 1:
                    raise ConsolidationError(f'{name}: maximum {MAX_ROWS:,} worksheet rows supported.')
                if any(c.data_type == 'f' for c in row):
                    raise ConsolidationError(f'{name}: formula found on row {number}. Save formulas as values before uploading.')
                values = tuple(c.value for c in row)
                if any(v is not None for v in values):
                    rows.append(values)
            return WorkbookData(headers, rows, size)
        finally:
            book.close()
    except ConsolidationError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, ValueError, IndexError, OSError, SyntaxError) as exc:
        raise ConsolidationError(f'{name}: could not read this workbook. Upload a valid, unencrypted .xlsx file.') from exc
